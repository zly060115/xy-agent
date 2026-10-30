"""文档分析准确性测试：造 10 份埋好标准答案的夹具文档，让 Agent 读，比对答案。

用法：.venv/Scripts/python.exe -u .workbuddy/doc_test.py

夹具放在 workspace/doc_fixtures/（read_file 只能读 workspace 内的文件，
路径按 workspace 相对路径传）。每例给一个问题 + 预期关键词，自动判定。

覆盖 10 种文档分析能力：
1. TXT 关键数字提取     2. DOCX 段落信息     3. DOCX 表格单元格
4. PDF 文本提取         5. PDF 跨页汇总      6. 图片 OCR（印刷文字）
7. 图片图表读数         8. 长文档章节检索    9. 双文档对比
10. 文档没有的信息，拒绝编造
"""
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")

from dotenv import load_dotenv

load_dotenv(ROOT / ".env")

from myagent.config import config

FIXDIR = config.workspace_dir / "doc_fixtures"
FIXDIR.mkdir(parents=True, exist_ok=True)

# ---------------- 生成夹具 ----------------

def make_txt(name, content):
    p = FIXDIR / name
    p.write_text(content, encoding="utf-8")
    return p.name


def make_docx(name, paragraphs, table=None):
    from docx import Document
    doc = Document()
    for t in paragraphs:
        doc.add_paragraph(t)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, cell in enumerate(row):
                t.rows[i].cells[j].text = cell
    p = FIXDIR / name
    doc.save(p)
    return p.name


def make_pdf(name, pages):
    from fpdf import FPDF
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_font("cn", "", "C:/Windows/Fonts/simhei.ttf")
    for text in pages:
        pdf.add_page()
        pdf.set_font("cn", size=14)
        pdf.multi_cell(0, 10, text)
    p = FIXDIR / name
    pdf.output(str(p))
    return p.name


def make_img_text(name, lines):
    from PIL import Image, ImageDraw, ImageFont
    font = ImageFont.truetype("C:/Windows/Fonts/simhei.ttf", 32)
    img = Image.new("RGB", (760, 70 + 56 * len(lines)), "white")
    d = ImageDraw.Draw(img)
    for i, t in enumerate(lines):
        d.text((40, 30 + 56 * i), t, fill="black", font=font)
    p = FIXDIR / name
    img.save(p)
    return p.name


def make_img_chart(name):
    """手画柱状图：4 根柱子，高度带数值标签，测 OCR 读数。"""
    from PIL import Image, ImageDraw, ImageFont
    font = ImageFont.truetype("C:/Windows/Fonts/simhei.ttf", 28)
    W, H = 760, 520
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    d.text((200, 16), "各季度销量（件）", fill="black", font=font)
    data = [("Q1", 350), ("Q2", 820), ("Q3", 640), ("Q4", 910)]
    maxv = 1000
    x0, bw, gap, base = 90, 110, 70, 440
    for i, (label, v) in enumerate(data):
        x = x0 + i * (bw + gap)
        h = int(340 * v / maxv)
        d.rectangle([x, base - h, x + bw, base], fill="#4472c4")
        d.text((x + 22, base - h - 40), str(v), fill="black", font=font)
        d.text((x + 24, base + 12), label, fill="black", font=font)
    d.line([60, base, 740, base], fill="black", width=3)
    p = FIXDIR / name
    img.save(p)
    return p.name


LONG_MD = """# 机房运维手册（节选）

## 1. 温度管理
机房温度应保持在 18 至 27 摄氏度之间。

## 2. 消防要求
机房内应配备气体灭火器，禁止使用干粉灭火器。

## 3. 供电
双路市电接入，UPS 续航不低于 30 分钟。

## 4. 网络布线
强弱电分离，线缆每隔 1.5 米固定一次。

## 5. 门禁
进出机房需刷卡登记，访客须由值班工程师陪同。

## 6. 湿度管理
机房相对湿度应保持在 45% 到 60% 之间。

## 7. 卫生
每周清洁一次地板，禁止堆放纸箱。

## 8. 应急联系
值班电话 6888，夜间联系人王强。
"""

f1 = make_txt("01_生产报告.txt",
              "Q3 生产月报：本月总产量 1200 台，一次良率 97.6%，返修 29 台，"
              "出货 1180 台，产线稼动率 91%。")
f2 = make_docx("02_采购合同.docx",
               ["设备采购合同（摘要）",
                "乙方向甲方采购检测设备，设备单价 3500 元，采购数量 42 台，"
                "合同总价 147000 元，交货期为签约后 30 天。",
                "设备质保期两年。"])
f3 = make_docx("03_销售表.docx",
               ["2026 年分季度销售额（万元）"],
               table=[["季度", "销售额（万元）"],
                      ["Q1", "85"], ["Q2", "92"], ["Q3", "108"], ["Q4", "131"]])
f4 = make_pdf("04_实验报告.pdf",
              ["模型评测实验报告：共使用样本 500 条，模型准确率 93.2%，"
               "单轮推理平均耗时 210 毫秒，训练总耗时 4.5 小时。"])
f5 = make_pdf("05_项目立项书.pdf",
              ["项目立项书（第一页）：项目名称为智能巡检平台，项目负责人张伟，"
               "立项日期 2026 年 9 月。",
               "项目立项书（第二页）：项目总预算 26 万元，其中设备采购 15 万元，"
               "人工 8 万元，其他 3 万元。"])
f6 = make_img_text("06_保修卡.png",
                   ["产品保修卡", "产品编号：A-2026-88",
                    "保修期：24 个月", "客服电话：400-800-1234"])
f7 = make_img_chart("07_柱状图.png")
f8 = make_txt("08_运维手册.md", LONG_MD)
f9a = make_txt("09a_甲厂产量.txt", "甲厂 10 月产量：1500 台，合格率 98.1%。")
f9b = make_txt("09b_乙厂产量.txt", "乙厂 10 月产量：1300 台，合格率 98.9%。")

if "--fixtures-only" in sys.argv:
    for p in sorted(FIXDIR.iterdir()):
        print(p.name, p.stat().st_size)
    sys.exit(0)

# ---------------- 用例 ----------------

def need(*keys):
    return ("all", keys)

def any_of(*keys):
    return ("any", keys)

CASES = [
    ("TXT 提取",  f1, "读一下 doc_fixtures/01_生产报告.txt，一次良率是多少？", need("97.6")),
    ("DOCX 段落", f2, "读一下 doc_fixtures/02_采购合同.docx，合同总价是多少？", need("147000")),
    ("DOCX 表格", f3, "读一下 doc_fixtures/03_销售表.docx，Q3 的销售额是多少万元？", need("108")),
    ("PDF 文本",  f4, "读一下 doc_fixtures/04_实验报告.pdf，模型准确率是多少？", need("93.2")),
    ("PDF 跨页",  f5, "读一下 doc_fixtures/05_项目立项书.pdf，项目总预算是多少？负责人是谁？",
     need("26", "张伟")),
    ("图片 OCR",  f6, "读一下 doc_fixtures/06_保修卡.png，保修期是多久？", need("24")),
    ("图表读数",  f7, "读一下 doc_fixtures/07_柱状图.png，销量最高的季度是多少件？", need("910")),
    ("章节检索",  f8, "读一下 doc_fixtures/08_运维手册.md，机房湿度应该保持在什么范围？",
     need("45", "60")),
    ("双文档对比", None, "分别读一下 doc_fixtures/09a_甲厂产量.txt 和 09b_乙厂产量.txt，"
     "哪个厂产量更高？多少台？", need("甲", "1500")),
    ("拒绝编造",  f2, "读一下 doc_fixtures/02_采购合同.docx，这份合同的付款方式是什么？",
     any_of("没有", "未", "不包含", "并未", "没有提到", "没有约定")),
]

# ---------------- 跑 ----------------

import myagent.agent as agent_mod
from myagent.agent import Agent
from myagent.confirm import ConfirmGateway
from myagent.config import config
from myagent.llm import build_llm
from myagent.rag import RAGStore, build_embed_fn
from myagent.runtime import set_runtime
from myagent.safe_reader import SafeFileReader

llm = build_llm(config)
set_runtime(
    llm=llm,
    rag=RAGStore(embed_fn=build_embed_fn()),
    reader=SafeFileReader(vision=llm),
    gateway=ConfirmGateway(mode="auto_approve"),
)

passed = 0
for label, fixture, q, (mode, keys) in CASES:
    agent = Agent(llm, verbose=False, use_rag=False)
    t = time.time()
    try:
        ans = agent.run(q)
    except Exception as e:
        ans = f"<异常 {type(e).__name__}: {e}>"
    if mode == "all":
        ok = all(k in ans for k in keys)
    else:
        ok = any(k in ans for k in keys)
    passed += ok
    print(f"{'✅' if ok else '❌'} {label}  预期含 {keys}  {time.time()-t:.0f}s")
    print("   " + ans.replace("\n", " ")[:180])
    print()

print(f"通过 {passed}/{len(CASES)}")
