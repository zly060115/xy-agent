"""写入能力回归：工作区免授权、用户目录可写、越界一律拒绝。

跑法：.venv/Scripts/python.exe .workbuddy/write_test.py
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from myagent.config import config  # noqa: E402
from myagent.tools import (  # noqa: E402
    TOOLS,
    WORKSPACE,
    _resolve_write_path,
    _write_blocker,
    _write_roots,
    tool_write_file,
)

ok = 0
fail = 0
cleanup: list[Path] = []


def check(name: str, cond: bool, extra: str = ""):
    global ok, fail
    if cond:
        ok += 1
        print(f"  [PASS] {name}")
    else:
        fail += 1
        print(f"  [FAIL] {name}  {extra}")


print("=" * 60)
print("1. 可写范围")
roots = [str(r) for r in _write_roots()]
print("   可写根目录:", roots)
check("工作区在可写范围内", any("workspace" in r for r in roots))
check("桌面在可写范围内", any(r.lower().endswith("desktop") for r in roots))
check("可写目录不止工作区", len(_write_roots()) >= 2)

print("=" * 60)
print("2. 工作区相对路径写入（免授权）")
rel = "write_test/note.md"
r = tool_write_file(rel, "# 标题\n第一行内容\n")
p = WORKSPACE / rel
cleanup.append(p)
check("返回成功", r.startswith("✅"), r)
check("文件真的存在", p.exists())
check("内容正确", p.exists() and "第一行内容" in p.read_text(encoding="utf-8"))

print("=" * 60)
print("3. 桌面绝对路径写入 .md")
desk = next((d for d in config.write_dirs if d.name.lower() == "desktop"), None)
if desk is None:
    print("  [SKIP] 没找到桌面目录")
else:
    target = desk / "_xiaoman_write_test.md"
    cleanup.append(target)
    r = tool_write_file(str(target), "# 桌面测试\n由小满写入。\n")
    check("返回成功", r.startswith("✅"), r)
    check("报告里带「桌面」", "桌面" in r, r)
    check("文件落在桌面", target.exists())
    check("内容正确", target.exists() and "由小满写入" in target.read_text(encoding="utf-8"))

    print("=" * 60)
    print("4. 桌面写 .docx（真 Word 文档）")
    dx = desk / "_xiaoman_write_test.docx"
    cleanup.append(dx)
    md = "# 测试标题\n\n这是正文段落。\n\n## 小节\n\n- 要点一\n- 要点二\n"
    r = tool_write_file(str(dx), md)
    check("返回成功", r.startswith("✅"), r)
    check("文件落在桌面", dx.exists())
    try:
        from docx import Document

        doc = Document(str(dx))
        texts = [p.text for p in doc.paragraphs]
        joined = "\n".join(texts)
        check("是合法 docx 能被读回", len(texts) > 0)
        check("标题写入成功", "测试标题" in joined, joined[:80])
        check("列表项写入成功", "要点一" in joined and "要点二" in joined)
        check("Markdown 符号没被抄进正文", "##" not in joined and "- 要点" not in joined)
        styles = {p.style.name for p in doc.paragraphs}
        check("标题用了 Heading 样式", any(s.startswith("Heading") for s in styles), str(styles))
    except Exception as e:
        check("docx 解析", False, f"{type(e).__name__}: {e}")

print("=" * 60)
print("5. 越界路径一律拒绝")
bad_cases = [
    ("目录穿越", str(WORKSPACE / ".." / ".." / "Windows" / "evil.md")),
    ("系统目录", "C:/Windows/System32/evil.md"),
    ("程序目录", "C:/Program Files/evil.md"),
    ("用户根目录", str(Path.home() / "evil.md")),
    ("临时目录", "C:/Windows/Temp/evil.md"),
]
for name, path in bad_cases:
    why = _write_blocker(_resolve_write_path(path))
    check(f"{name} 被拒", why is not None, f"{path} 竟然可写")
    if why:
        print(f"         理由: {why}")

print("=" * 60)
print("6. 可执行/非文档扩展名一律拒绝")
if desk is not None:
    for ext in (".bat", ".exe", ".ps1", ".vbs", ".lnk", ".py", ""):
        target = desk / f"_xiaoman_bad{ext}"
        why = _write_blocker(_resolve_write_path(str(target)))
        check(f"桌面写 {ext or '无扩展名'} 被拒", why is not None, "竟然可写")
    # 确认没有被偷偷创建
    created = [p for p in desk.glob("_xiaoman_bad*")]
    check("坏扩展名没落盘", not created, str(created))

print("=" * 60)
print("7. 工具说明书让模型知道真实可写路径")
desc = TOOLS["write_file"]["spec"]["description"]
check("说明书提到桌面真实路径", str(desk) in desc if desk else True, desc[:120])
check("说明书提到 docx", ".docx" in desc)
check("说明书禁止让用户自己拷", "自己" in desc or "拷" in desc)

print("=" * 60)
print("8. 系统提示词")
from myagent.agent import Agent  # noqa: E402
from myagent.llm import UnavailableLLM  # noqa: E402

sysmsg = Agent(UnavailableLLM(), verbose=False).messages[0]["content"]
check("提示词给了可写目录", str(desk) in sysmsg if desk else True)
check("提示词禁止「我只能写在规定目录」", "我只能写在规定目录" in sysmsg)
check("提示词要求写完回报路径", "完整路径" in sysmsg)

# ---- 清理 ----
for f in cleanup:
    try:
        if f.exists():
            f.unlink()
    except OSError as e:
        print(f"  清理失败 {f}: {e}")
left = [p for p in (desk.glob("_xiaoman_write_test*") if desk else [])]
check("测试文件已清理", not left, str(left))

print("=" * 60)
print(f"通过 {ok} 项，失败 {fail} 项")
sys.exit(0 if fail == 0 else 1)
